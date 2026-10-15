import sys
import io
import csv
import threading
from pathlib import Path
import pytest
import openpyxl

sys.path.insert(0, str(Path(__file__).parents[1]))
from app_server import make_server
from books import Books
from interchange import generate_sample_template, bulk_import_masters


def test_sample_template_generation():
    # 1. Ledgers CSV
    csv_ledgers = generate_sample_template('ledgers', 'csv')
    reader = csv.reader(io.StringIO(csv_ledgers))
    headers = next(reader)
    assert 'Ledger Name' in headers
    assert 'Parent Group' in headers
    assert 'Opening Balance' in headers
    assert 'GSTIN' in headers
    sample_row = next(reader)
    assert len(sample_row) == len(headers)

    # 2. Items CSV
    csv_items = generate_sample_template('items', 'csv')
    reader_items = csv.reader(io.StringIO(csv_items))
    headers_items = next(reader_items)
    assert 'Item Name' in headers_items
    assert 'Base Unit' in headers_items
    assert 'GST Rate %' in headers_items
    assert 'Standard Purchase Rate' in headers_items
    assert 'Standard Selling Rate' in headers_items
    sample_item_row = next(reader_items)
    assert len(sample_item_row) == len(headers_items)

    # 3. Excel .xlsx templates
    xlsx_ledgers_bytes = generate_sample_template('ledgers', 'xlsx')
    wb_l = openpyxl.load_workbook(io.BytesIO(xlsx_ledgers_bytes))
    sheet_l = wb_l.active
    assert sheet_l.cell(row=1, column=1).value == 'Ledger Name'

    xlsx_items_bytes = generate_sample_template('items', 'xlsx')
    wb_i = openpyxl.load_workbook(io.BytesIO(xlsx_items_bytes))
    sheet_i = wb_i.active
    assert sheet_i.cell(row=1, column=1).value == 'Item Name'


def test_bulk_import_ledgers_and_items_csv(tmp_path):
    db = Books(str(tmp_path / 'books.sqlite3'))
    db.create_or_alter_company({'name': 'Bulk Co', 'starting_from': '20260401', 'ending_at': '20270331'})

    # Import ledgers CSV
    ledgers_csv = (
        "Ledger Name,Parent Group,Opening Balance,Dr/Cr,GSTIN,State,Address,Pincode,Credit Days\n"
        "Apex Corp,Sundry Debtors,15000,Dr,19AAAAA1111A1Z1,West Bengal,Park Street,700016,30\n"
        "Beta Traders,Sundry Creditors,8000,Cr,19BBBBB2222B1Z2,West Bengal,Burrabazar,700007,15\n"
    )
    result = bulk_import_masters(db, 'Bulk Co', 'ledgers', ledgers_csv, 'csv')
    assert result['created'] == 2
    company = db.get_company('Bulk Co')
    assert 'Apex Corp' in company['ledgers']
    assert company['ledgers']['Apex Corp']['opening_balance'] == '15000.00'
    assert company['ledgers']['Apex Corp']['state'] == 'West Bengal'
    assert 'Beta Traders' in company['ledgers']

    # Import items CSV
    items_csv = (
        "Item Name,Stock Group,Base Unit,Alt Unit,Conversion Rate,HSN Code,GST Rate %,Standard Purchase Rate,Standard Selling Rate,Opening Qty,Opening Rate\n"
        "Steel Rod 10MM,Primary,pcs,,,7214,18,50.00,62.00,100,50.00\n"
        "Binding Wire,Primary,kgs,,,7217,18,40.00,48.00,200,40.00\n"
    )
    result_items = bulk_import_masters(db, 'Bulk Co', 'items', items_csv, 'csv')
    assert result_items['created'] == 2
    company = db.get_company('Bulk Co')
    assert 'Steel Rod 10MM' in company['stock_items']
    rod = company['stock_items']['Steel Rod 10MM']
    assert rod['standard_purchase_rate'] == '50.00'
    assert rod['standard_selling_rate'] == '62.00'
    assert rod['gst_rate'] == '18'


def test_bulk_import_atomic_rollback_on_invalid_group(tmp_path):
    db = Books(str(tmp_path / 'books.sqlite3'))
    db.create_or_alter_company({'name': 'Rollback Co', 'starting_from': '20260401', 'ending_at': '20270331'})

    # 1 valid row, 1 row with non-existent parent group
    bad_csv = (
        "Ledger Name,Parent Group,Opening Balance,Dr/Cr,GSTIN,State,Address,Pincode,Credit Days\n"
        "Valid Ledger,Sundry Debtors,1000,Dr,,,,,\n"
        "Broken Ledger,NonExistentGroup123,2000,Dr,,,,,\n"
    )
    with pytest.raises(ValueError, match="Unknown ledger group"):
        bulk_import_masters(db, 'Rollback Co', 'ledgers', bad_csv, 'csv')

    # Atomic rollback check: Valid Ledger must NOT be present
    company = db.get_company('Rollback Co')
    assert 'Valid Ledger' not in company['ledgers']
    assert 'Broken Ledger' not in company['ledgers']


def test_bulk_import_api_endpoints(tmp_path):
    server = make_server(0, str(tmp_path / 'books.sqlite3'))
    db = server.db
    db.create_or_alter_company({'name': 'API Co', 'starting_from': '20260401', 'ending_at': '20270331'})

    import urllib.request
    import json

    base_url = f'http://127.0.0.1:{server.server_port}'
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()

    try:
        # 1. Download sample template via GET
        req = urllib.request.Request(f'{base_url}/api/sample_template?type=ledgers&format=csv')
        with urllib.request.urlopen(req) as resp:
            assert resp.status == 200
            content = resp.read().decode('utf-8')
            assert 'Ledger Name' in content

        # 2. Upload bulk import via POST
        post_data = json.dumps({
            'company': 'API Co',
            'type': 'ledgers',
            'format': 'csv',
            'content': "Ledger Name,Parent Group,Opening Balance,Dr/Cr\nDirect API Party,Sundry Debtors,500,Dr\n"
        }).encode('utf-8')
        post_req = urllib.request.Request(f'{base_url}/api/bulk_import', data=post_data, headers={'Content-Type': 'application/json'})
        with urllib.request.urlopen(post_req) as resp:
            assert resp.status == 200
            res = json.loads(resp.read().decode('utf-8'))
            assert res['status'] == 'success'
            assert res['created'] == 1

        company = db.get_company('API Co')
        assert 'Direct API Party' in company['ledgers']
    finally:
        server.shutdown()
        thread.join(timeout=3)
        server.server_close()
