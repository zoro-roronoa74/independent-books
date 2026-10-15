"""Validated portable backups and Tally XML master/voucher interchange."""
import copy
import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from books import Books, GROUP_PARENTS, KINDS, VOUCHER_TYPES, decimal, entries, date_key

TAGS={'GROUP':'group','LEDGER':'ledger','STOCKITEM':'stock_item','UNIT':'unit','GODOWN':'godown',
      'VOUCHERTYPE':'voucher_type','STOCKGROUP':'stock_group','STOCKCATEGORY':'stock_category',
      'COSTCATEGORY':'cost_category','COSTCENTRE':'cost_centre','CURRENCY':'currency'}


def local_tag(tag):
    """Return an XML local name while ignoring an optional namespace."""
    return str(tag).rsplit('}', 1)[-1]


def _findtext(element, name, default=''):
    for node in element.iter():
        if local_tag(node.tag) == name:
            return node.text or default
    return default


def _findall(element, name):
    return [node for node in element.iter() if local_tag(node.tag) == name]


_XML_CHAR_REF=re.compile(r'&#(?:x[0-9A-Fa-f]+|[0-9]+);')
_XML_ILLEGAL_RAW=re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]')


def _xml_safe_text(text):
    """Remove illegal XML 1.0 control characters and references emitted by Tally."""
    if isinstance(text, bytes):
        for enc in ('utf-16', 'utf-8-sig', 'utf-8', 'latin-1'):
            try:
                text = text.decode(enc)
                break
            except UnicodeDecodeError:
                pass
    clean_str = _XML_ILLEGAL_RAW.sub('', str(text).lstrip('\ufeff'))
    def replace(match):
        token=match.group(0)
        value=int(token[3:-1],16) if token[2].lower() == 'x' else int(token[2:-1])
        valid=value in (0x9,0xA,0xD) or 0x20 <= value <= 0xD7FF or 0xE000 <= value <= 0xFFFD or 0x10000 <= value <= 0x10FFFF
        return token if valid else ''
    return _XML_CHAR_REF.sub(replace, clean_str)


def xml_dict(element):
    result={}
    for node in element:
        tag=node.tag if str(node.tag).startswith('{') else local_tag(node.tag)
        value=xml_dict(node) if len(node) else (node.text or '')
        if tag.endswith('.LIST'):
            result.setdefault(tag,[]).append(value)
        elif tag in result:
            if not isinstance(result[tag],list): result[tag]=[result[tag]]
            result[tag].append(value)
        else: result[tag]=value
    result.update({'@'+k:v for k,v in element.attrib.items()})
    return result


def _clean_empty_lists(record):
    """Normalize empty Tally LIST elements to empty arrays, retaining raw XML."""
    if isinstance(record, dict):
        for key, value in list(record.items()):
            if key.endswith('.LIST') and isinstance(value, list):
                record[key]=[item for item in value if not isinstance(item,str) or item.strip()]
                value=record[key]
            _clean_empty_lists(value)
    elif isinstance(record, list):
        for value in record: _clean_empty_lists(value)
    return record


def _normalize_voucher_allocations(record):
    """Promote item accounting allocations to ledger entries for the books engine."""
    _clean_empty_lists(record)
    ledger_entries=entries(record)
    existing_names={entry.get('LEDGERNAME') for entry in ledger_entries if isinstance(entry,dict)}
    promoted=[]
    for key in ('ALLINVENTORYENTRIES.LIST','INVENTORYENTRIES.LIST'):
        for item in entries(record,key):
            if not isinstance(item,dict): continue
            for allocation in entries(item,'ACCOUNTINGALLOCATIONS.LIST'):
                if isinstance(allocation,dict) and allocation.get('LEDGERNAME') not in existing_names:
                    promoted.append(copy.deepcopy(allocation))
    if promoted:
        record['ALLLEDGERENTRIES.LIST']=[*ledger_entries,*promoted]
    return record


def parse_tally_bundle(xmls):
    """Parse one or more Tally XML documents into a single import plan."""
    if not isinstance(xmls, list) or not xmls:
        raise ValueError('At least one Tally XML file is required')
    messages=[]; skipped=[]; sources=[]; company_name=''; company_node=None; company_names=set()
    for index, source in enumerate(xmls):
        if isinstance(source, dict):
            name=str(source.get('name') or f'xml-{index + 1}.xml')
            text=source.get('text', '')
        else:
            name=f'xml-{index + 1}.xml'; text=source or ''
        if isinstance(text, bytes):
            for enc in ('utf-16', 'utf-8-sig', 'utf-8', 'latin-1'):
                try:
                    text = text.decode(enc)
                    break
                except UnicodeDecodeError:
                    pass
        text=str(text or '')
        if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
            raise ValueError(f'{name}: XML document types and entities are not supported')
        try:
            root=ET.fromstring(_xml_safe_text(text))
        except ET.ParseError as exc:
            raise ValueError(f'{name}: invalid XML: {exc}') from exc
        sources.append({'name':name,'text':text})
        current=_findtext(root, 'SVCURRENTCOMPANY').strip()
        if current: company_names.add(current)
        if current and not company_name: company_name=current
        file_messages=[]
        for tally_message in _findall(root, 'TALLYMESSAGE'):
            children=list(tally_message)
            for node in children:
                entity=local_tag(node.tag)
                file_messages.append({'entity':entity,'node':node,'file_index':index,
                    'file_name':name,'xml':ET.tostring(tally_message,encoding='unicode'),
                    'name':node.get('NAME') or _findtext(node,'NAME').strip(),
                    'stable_id':node.get('REMOTEID') or _findtext(node,'GUID').strip()})
                if entity == 'COMPANY':
                    company_node=node
                    remote_name=_findtext(node,'REMOTECMPNAME').strip()
                    if remote_name:
                        company_names.add(remote_name)
                        if not company_name: company_name=remote_name
        if not file_messages:
            skipped.append({'name':name,'reason':'No TALLYMESSAGE records found'})
        messages.extend(file_messages)
    if len(company_names) > 1:
        raise ValueError('XML bundle contains different companies: '+', '.join(sorted(company_names)))
    if not company_name and company_node is not None:
        company_name=(company_node.get('NAME') or _findtext(company_node,'NAME')).strip()
    return {'company':{'name':company_name,'node':company_node},'messages':messages,
            'skipped':skipped,'sources':sources}


def validate_backup(data):
    if not isinstance(data,dict) or data.get('version')!=1 or not isinstance(data.get('companies'),dict): raise ValueError('Unsupported backup schema')
    if not isinstance(data.get('active_company'),str) or not isinstance(data.get('revision'),int): raise ValueError('Backup lacks required application metadata')
    companies=data['companies']
    if (companies and data['active_company'] not in companies) or (not companies and data['active_company']): raise ValueError('Invalid active company in backup')
    identities=set()
    for name,c in companies.items():
        if not isinstance(c,dict) or c.get('name')!=name or not name: raise ValueError('Backup company identity does not match')
        if any(not isinstance(c.get(key),str) or not c[key] for key in ('guid','company_number','starting_from','ending_at')): raise ValueError('Backup lacks required company metadata')
        if c['guid'] in identities: raise ValueError('Duplicate company GUID in backup')
        identities.add(c['guid'])
        if date_key(c['ending_at'])<date_key(c['starting_from']): raise ValueError('Invalid financial period in backup')
        if date_key(c.get('books_from') or c['starting_from'])<date_key(c['starting_from']): raise ValueError('Invalid books beginning date in backup')
        for key in ('groups','voucher_types','vouchers','units','godowns','stock_groups','stock_categories','cost_categories','cost_centres','currencies'):
            if not isinstance(c.get(key),list): raise ValueError('Invalid backup collection: '+key)
        for key in ('ledgers','stock_items','group_records'):
            if not isinstance(c.get(key),dict) or any(not isinstance(value,dict) for value in c[key].values()): raise ValueError('Invalid backup master collection: '+key)
        for field in KINDS.values():
            if field not in ('ledgers','stock_items','currencies'):
                if any(not isinstance(n,str) or not n for n in c[field]) or len(set(c[field]))!=len(c[field]): raise ValueError('Invalid or duplicate master name in backup')
        for key in ('ledgers','stock_items'):
            if any(record.get('name')!=name_ for name_,record in c[key].items()): raise ValueError('Backup master identity does not match')
        for ledger in c['ledgers'].values():
            if ledger.get('parent') not in c['groups']: raise ValueError('Unknown ledger group in backup')
            decimal(ledger.get('signed_opening'),'.01')
        for item in c['stock_items'].values():
            if item.get('base_units') and item['base_units'] not in c['units']: raise ValueError('Unknown stock unit in backup')
            if decimal(item.get('opening_qty'))<0 or decimal(item.get('opening_balance'))<0: raise ValueError('Negative opening stock in backup')
        for v in c['vouchers']:
            if not isinstance(v,dict) or not isinstance(v.get('GUID'),str) or not v['GUID'] or not isinstance(v.get('MASTERID'),str): raise ValueError('Backup lacks required voucher identity')
            if not v['MASTERID'].startswith('v-') or not v['MASTERID'][2:].isdigit(): raise ValueError('Invalid voucher ID in backup')


def merged_master_xml(previous,node):
    full=ET.fromstring(previous['raw_xml']) if previous.get('raw_xml') else ET.Element(node.tag)
    full.attrib.update(node.attrib)
    for tag in set(child.tag for child in node):
        for old in list(full):
            if old.tag==tag: full.remove(old)
        for child in node:
            if child.tag==tag: full.append(copy.deepcopy(child))
    return full


def _candidate(books,payload):
    candidate=Books(':memory:');candidate.data=copy.deepcopy(books.data)
    counts={'companies':0,'masters':0,'vouchers':0,'updated':0,'warnings':[],
            'skipped':[],'errors':[]}
    try:
        if payload.get('format')=='backup':
            data=payload['data']
            if isinstance(data,str): data=json.loads(data)
            validate_backup(data)
            candidate.data=copy.deepcopy(data)
            for name,c in candidate.data['companies'].items():
                if c.get('name')!=name: raise ValueError('Backup company identity does not match')
                for group in c['groups']: candidate.root_group(c,group)
                candidate.view(name)
                ids=set()
                guids=set()
                numbers=set()
                for v in c['vouchers']:
                    candidate._validate_voucher(c,v)
                    if v['MASTERID'] in ids: raise ValueError('Duplicate voucher ID in backup')
                    ids.add(v['MASTERID'])
                    if v.get('GUID') in guids: raise ValueError('Duplicate voucher GUID in backup')
                    guids.add(v.get('GUID'))
                    number=(v.get('VOUCHERTYPENAME'),str(v.get('VOUCHERNUMBER','')))
                    if not number[1] or number in numbers: raise ValueError('Invalid or duplicate voucher number in backup')
                    numbers.add(number)
                c['next_voucher_id']=max(int(c.get('next_voucher_id',1)),max((int(str(x).split('-')[-1])+1 for x in ids),default=1))
                counts['masters']+=len(c['ledgers'])+len(c['stock_items'])
                counts['vouchers']+=len(c['vouchers'])
            if data.get('active_company') and data['active_company'] not in data['companies']: raise ValueError('Invalid active company in backup')
            counts['companies']=len(data['companies'])
            return candidate,counts
        node_context={}
        if payload.get('xmls') is not None:
            bundle=parse_tally_bundle(payload['xmls'])
            text='\n'.join(source['text'] for source in bundle['sources'])
            source_texts=[source['text'] for source in bundle['sources']]
            nodes=[message['node'] for message in bundle['messages']]
            node_context={id(message['node']):message for message in bundle['messages']}
            name=str(payload.get('company') or bundle['company']['name'] or '').strip()
            company_node=bundle['company']['node']
            counts['skipped'].extend(bundle['skipped'])
            if not nodes: raise ValueError('No importable Tally messages were found in the XML bundle')
        else:
            text=payload.get('xml','')
            if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper(): raise ValueError('XML document types and entities are not supported')
            root=ET.fromstring(_xml_safe_text(text))
            nodes=[]
            for tally_message in _findall(root,'TALLYMESSAGE'):
                nodes.extend(list(tally_message))
            source_texts=[text]
            name=str(payload.get('company') or _findtext(root,'SVCURRENTCOMPANY') or '').strip()
            company_nodes=[node for node in nodes if local_tag(node.tag)=='COMPANY']
            company_node=company_nodes[0] if company_nodes else None
            if not name and company_node is not None:
                name=company_node.get('NAME') or _findtext(company_node,'NAME') or ''
        if not name: raise ValueError('Select a company or provide the company name for import')
        if name not in candidate.data['companies']:
            metadata=json.loads(_findtext(company_node,'BOOKSMETADATA','{}')) if company_node is not None else {}
            if not isinstance(metadata,dict): raise ValueError('Invalid company metadata in XML')
            start_date = str(payload.get('starting_from') or (_findtext(company_node,'STARTINGFROM') if company_node is not None else '') or '').strip()
            if not start_date:
                for node in nodes:
                    s = _findtext(node, 'STARTINGFROM') or _findtext(node, 'APPLICABLEFROM') or _findtext(node, 'DATE')
                    if s:
                        try:
                            start_date = date_key(s)
                            break
                        except ValueError:
                            pass
            if not start_date:
                import datetime as _dt
                today = _dt.date.today()
                fy_year = today.year if today.month >= 4 else today.year - 1
                start_date = f'{fy_year}0401'
            else:
                try:
                    start_date = date_key(start_date)
                except ValueError:
                    pass
            end_date = str(payload.get('ending_at') or (_findtext(company_node,'ENDINGAT') if company_node is not None else '') or '').strip()
            if not end_date:
                try:
                    sk = date_key(start_date)
                    end_date = f'{int(sk[:4]) + 1}0331'
                except ValueError:
                    end_date = '20270331'
            else:
                try:
                    end_date = date_key(end_date)
                except ValueError:
                    pass
            books_from = str(_findtext(company_node,'BOOKSFROM') if company_node is not None else '').strip() or start_date
            cmp_state = payload.get('state') or metadata.get('state') or (_findtext(company_node,'REMOTECMPSTATE') if company_node is not None else '')
            cmp_fields = {**metadata,'name':name,
                'starting_from':start_date,
                'ending_at':end_date,
                'books_from':books_from,
                'currency':payload.get('currency',metadata.get('currency',''))}
            if cmp_state: cmp_fields['state'] = cmp_state
            candidate.create_or_alter_company(cmp_fields)
            if metadata.get('name')==name and metadata.get('guid'):
                if any(c.get('guid')==metadata['guid'] for n,c in candidate.data['companies'].items() if n!=name): raise ValueError('Company GUID already exists')
                candidate.get_company(name)['guid']=metadata['guid']
            counts['companies']=1
        c=candidate.get_company(name)
        # Parent groups precede their children; ledgers/items follow units/groups.
        master_nodes=[n for n in nodes if local_tag(n.tag) in TAGS]
        order={'group':0,'unit':1,'stock_group':2,'stock_category':3,'godown':4,'ledger':5,'stock_item':6}
        master_nodes.sort(key=lambda n:order.get(TAGS[local_tag(n.tag)],2))
        group_nodes=[n for n in master_nodes if local_tag(n.tag)=='GROUP']
        for n in group_nodes:
            nm=n.get('NAME') or _findtext(n,'NAME')
            if nm and nm not in c['groups']: c['groups'].append(nm)
        with candidate.transaction():
            for node in master_nodes:
                action=node.get('ACTION','Create')
                if action not in ('Create','Alter',''): raise ValueError('Unsupported master action: '+action)
                tag=local_tag(node.tag);kind=TAGS[tag];raw=xml_dict(node)
                name_=node.get('NAME') or _findtext(node,'NAME') or _findtext(node,'ORIGINALNAME')
                if not name_: raise ValueError('Unnamed '+node.tag+' in import')
                old_name=_findtext(node,'ORIGINALNAME') or name_
                previous=(c[KINDS[kind]].get(old_name,{}) if isinstance(c[KINDS[kind]],dict) else c.get(kind+'_records',{}).get(old_name,{}))
                if action=='Alter' and old_name not in c[KINDS[kind]]: raise ValueError('Master to alter does not exist: '+old_name)
                full=merged_master_xml(previous,node) if action=='Alter' else node
                record={'name':name_,'old_name':old_name,'raw_xml':ET.tostring(full,encoding='unicode'),'raw_fields':xml_dict(full)}
                if node.find('PARENT') is not None or action!='Alter':
                    parent=_findtext(node,'PARENT') or 'Primary'
                    if kind=='ledger' and parent=='Primary': parent='Capital Account'
                    record['parent']=parent
                if kind=='group':
                    classification=node.get('BOOKSCLASSIFICATION') or node.get('RESERVEDNAME') or _findtext(node,'RESERVEDNAME')
                    if classification:
                        aliases={str(group).lower():group for group in GROUP_PARENTS}
                        canonical=aliases.get(str(classification).strip().lower())
                        if canonical: record['system_group']=canonical
                if kind=='ledger':
                    if node.find('OPENINGBALANCE') is not None or action!='Alter': record.update(opening_balance=str(abs(decimal(_findtext(node,'OPENINGBALANCE',0),'0.01'))),drcr='Dr' if decimal(_findtext(node,'OPENINGBALANCE',0))<=0 else 'Cr')
                    for tag,key in (('PARTYGSTIN','gstin'),('ADDRESS','address'),('ISBILLWISEON','is_billwise')):
                        if _findall(node,tag): record[key]=_findtext(node,tag,'')
                elif kind=='stock_item':
                    for tag,key in (('BASEUNITS','base_units'),('OPENINGBALANCE','opening_qty'),('OPENINGVALUE','opening_balance'),('OPENINGRATE','opening_rate'),('CATEGORY','category')):
                        if _findall(node,tag):
                            value=_findtext(node,tag,'')
                            if key=='base_units' and value.strip() in ('Not Applicable', 'None', ''): value = ''
                            record[key]=str(decimal(value.split('/')[0])) if tag.startswith('OPENING') else value
                elif kind=='unit':
                    for tag,key in (('UNITTYPE','type'),('BASEUNIT','base_unit'),('UNITFACTOR','factor'),('DECIMALPLACES','decimal_places')):
                        if _findall(node,tag):
                            record[key]=_findtext(node,tag,'')
                elif kind=='voucher_type' and (node.find('PARENT') is not None or action!='Alter'):
                    base_type=_findtext(node,'PARENT') or name_
                    record['base_type']=base_type if base_type in VOUCHER_TYPES else (name_ if name_ in VOUCHER_TYPES else 'Journal')
                candidate.save_master(name,kind,record);counts['masters']+=1
        for node in (n for n in nodes if local_tag(n.tag)=='VOUCHER'):
            if node.get('ACTION','Create') not in ('Create','Alter',''): raise ValueError('Unsupported voucher action: '+node.get('ACTION'))
            record=_normalize_voucher_allocations(xml_dict(node));record['raw_xml']=ET.tostring(node,encoding='unicode')
            record['VOUCHERTYPENAME']=record.get('VOUCHERTYPENAME') or node.get('VCHTYPE','')
            guid=record.get('GUID') or node.get('REMOTEID') or node.get('VCHKEY')
            if guid: record['GUID']=guid
            old=next((v for v in c['vouchers'] if guid and v.get('GUID')==guid),None)
            if old:
                record.update(MASTERID=old['MASTERID'],**{'@ACTION':'Alter'});counts['updated']+=1
            else:
                if node.get('ACTION')=='Alter': raise ValueError('XML alteration requires an existing voucher GUID')
                record.pop('MASTERID',None);record['@ACTION']='Create'
            try:
                candidate.post_voucher(name,record)
            except ValueError as exc:
                context=node_context.get(id(node),{})
                label=context.get('file_name','XML')+': '+context.get('entity','VOUCHER')
                source_name=context.get('name') or record.get('VOUCHERNUMBER') or record.get('GUID') or ''
                raise ValueError(f'{label} {source_name}: {exc}') from exc
            counts['vouchers']+=1
        c.setdefault('source_imports',[]).append({'xml':text,'xmls':source_texts,
            'masters':counts['masters'],'vouchers':counts['vouchers'],'skipped':counts['skipped']})
        counts['warnings'].append('Original XML and unrecognised fields are retained. GST filing, tracking allocations and specialised valuation require separate parity checks.')
        return candidate,counts
    except Exception:
        candidate.connection.close();raise


def preview_import(books,payload):
    candidate,counts=_candidate(books,payload)
    candidate.connection.close()
    return {'status':'success',**counts}


def apply_import(books,payload):
    with books.transaction():
        revision=books.data['revision']
        candidate,counts=_candidate(books,payload)
        try:
            books.data=copy.deepcopy(candidate.data);books.data['revision']=revision
        finally: candidate.connection.close()
    return {'status':'success',**counts}


def _append(parent,key,value):
    if key.startswith('@'):
        parent.set(key[1:],str(value));return
    if key=='raw_xml' or key=='raw_fields': return
    if isinstance(value,list):
        for item in value: _append(parent,key,item)
        return
    node=ET.SubElement(parent,key)
    if isinstance(value,dict):
        for k,v in value.items(): _append(node,k,v)
    else: node.text=str(value)


def export_xml(company):
    root=ET.Element('ENVELOPE');header=ET.SubElement(root,'HEADER');ET.SubElement(header,'TALLYREQUEST').text='Import Data'
    body=ET.SubElement(root,'BODY');imp=ET.SubElement(body,'IMPORTDATA');desc=ET.SubElement(imp,'REQUESTDESC')
    ET.SubElement(desc,'REPORTNAME').text='All Masters';static=ET.SubElement(desc,'STATICVARIABLES');ET.SubElement(static,'SVCURRENTCOMPANY').text=company['name']
    data=ET.SubElement(imp,'REQUESTDATA')
    cm=ET.SubElement(ET.SubElement(data,'TALLYMESSAGE'),'COMPANY',NAME=company['name'])
    ET.SubElement(cm,'STARTINGFROM').text=company['starting_from']
    ET.SubElement(cm,'ENDINGAT').text=company['ending_at']
    ET.SubElement(cm,'BOOKSFROM').text=company.get('books_from',company['starting_from'])
    omitted=set(KINDS.values())|{'vouchers','source_imports','next_voucher_id','bills_receivable','bills_payable'}
    ET.SubElement(cm,'BOOKSMETADATA').text=json.dumps({k:v for k,v in company.items() if k not in omitted and not k.endswith('_records')},ensure_ascii=False)
    for tag,kind in TAGS.items():
        values=company.get(KINDS[kind],{})
        if isinstance(values,dict): records=values.values()
        else: records=[company.get(kind+'_records',{}).get(v,{'name':v}) if isinstance(v,str) else v for v in values]
        for record in records:
            name=record.get('name') or record.get('symbol')
            if not name: continue
            element=ET.fromstring(record['raw_xml']) if record.get('raw_xml') else ET.Element(tag)
            element.set('NAME',name);element.set('ACTION','Create')
            if kind=='group' and record.get('system_group'): element.set('BOOKSCLASSIFICATION',record['system_group'])
            mapped={'PARENT':record.get('parent','Primary')}
            if kind=='voucher_type': mapped['PARENT']=record.get('base_type',name)
            if kind=='ledger': mapped.update(OPENINGBALANCE=str(-decimal(record.get('signed_opening',0),'0.01')),PARTYGSTIN=record.get('gstin',''))
            if kind=='ledger': mapped.update(ISBILLWISEON=record.get('is_billwise','No'),ADDRESS=record.get('address',''))
            if kind=='stock_item': mapped.update(BASEUNITS=record.get('base_units',''),OPENINGBALANCE=str(record.get('opening_qty',0)),OPENINGVALUE=str(record.get('opening_balance',0)),OPENINGRATE=str(record.get('opening_rate',0)))
            if kind=='stock_item': mapped['CATEGORY']=record.get('category','Primary')
            if kind=='unit': mapped.update(UNITTYPE=record.get('type','quantity'),BASEUNIT=record.get('base_unit',name),UNITFACTOR=str(record.get('factor','1')),DECIMALPLACES=str(record.get('decimal_places','')))
            for key,value in mapped.items():
                node=element.find(key)
                if node is None: node=ET.SubElement(element,key)
                node.text=str(value)
            ET.SubElement(data,'TALLYMESSAGE').append(element)
    for v in company['vouchers']:
        node=ET.SubElement(ET.SubElement(data,'TALLYMESSAGE'),'VOUCHER',VCHTYPE=v['VOUCHERTYPENAME'],ACTION='Create')
        for key,value in v.items():
            if key in ('MASTERID','raw_xml','@ACTION','@VCHTYPE'): continue
            _append(node,key,value)
    return ET.tostring(root,encoding='unicode')


LEDGER_TEMPLATE_HEADERS = [
    'Ledger Name', 'Parent Group', 'Opening Balance', 'Dr/Cr',
    'GSTIN', 'State', 'Address', 'Pincode', 'Credit Days'
]
LEDGER_SAMPLE_ROW = [
    'ABC Enterprises', 'Sundry Debtors', '10000.00', 'Dr',
    '19AAAAA0000A1Z5', 'West Bengal', '12 Park Street', '700016', '30'
]

ITEM_TEMPLATE_HEADERS = [
    'Item Name', 'Stock Group', 'Base Unit', 'Alt Unit', 'Conversion Rate',
    'HSN Code', 'GST Rate %', 'Standard Purchase Rate', 'Standard Selling Rate',
    'Opening Qty', 'Opening Rate'
]
ITEM_SAMPLE_ROW = [
    'TMT Bar 12MM', 'Primary', 'pcs', '', '',
    '7214', '18', '50.00', '65.00', '100', '50.00'
]


def _generate_pure_xlsx(headers, sample, title='Sheet1'):
    import zipfile
    import xml.sax.saxutils as sax
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>'
        ))
        z.writestr('_rels/.rels', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>'
        ))
        z.writestr('xl/_rels/workbook.xml.rels', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '</Relationships>'
        ))
        z.writestr('xl/workbook.xml', (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<sheets><sheet name="{sax.escape(title)}" sheetId="1" r:id="rId1"/></sheets>'
            '</workbook>'
        ))
        def _col_name(n):
            name = ''
            while n > 0:
                n, rem = divmod(n - 1, 26)
                name = chr(65 + rem) + name
            return name
        sheet_rows = []
        c_xml = [f'<c r="{_col_name(i)}1" t="inlineStr"><is><t>{sax.escape(str(h))}</t></is></c>' for i, h in enumerate(headers, 1)]
        sheet_rows.append(f'<row r="1">{"".join(c_xml)}</row>')
        c_xml = [f'<c r="{_col_name(i)}2" t="inlineStr"><is><t>{sax.escape(str(s))}</t></is></c>' for i, s in enumerate(sample, 1)]
        sheet_rows.append(f'<row r="2">{"".join(c_xml)}</row>')
        sheet_content = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<sheetData>{"".join(sheet_rows)}</sheetData>'
            '</worksheet>'
        )
        z.writestr('xl/worksheets/sheet1.xml', sheet_content)
    return buf.getvalue()


def generate_sample_template(template_type='ledgers', file_format='csv'):
    headers = LEDGER_TEMPLATE_HEADERS if template_type == 'ledgers' else ITEM_TEMPLATE_HEADERS
    sample = LEDGER_SAMPLE_ROW if template_type == 'ledgers' else ITEM_SAMPLE_ROW

    if file_format.lower() == 'csv':
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(headers)
        writer.writerow(sample)
        return out.getvalue()
    elif file_format.lower() in ('xlsx', 'excel'):
        title = 'Sample ' + template_type.capitalize()
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = title
            ws.append(headers)
            ws.append(sample)
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = Font(bold=True, color='1D4B2B')
                cell.fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()
        except Exception:
            return _generate_pure_xlsx(headers, sample, title=title)
    else:
        raise ValueError('Unsupported format: ' + str(file_format))


def bulk_import_masters(db, company, master_type, content, file_format='csv'):
    rows = []
    if file_format.lower() == 'csv':
        text = content.decode('utf-8') if isinstance(content, bytes) else str(content)
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif file_format.lower() in ('xlsx', 'excel'):
        import openpyxl
        data_bytes = content if isinstance(content, bytes) else content.encode('latin1')
        wb = openpyxl.load_workbook(io.BytesIO(data_bytes), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if all_rows:
            headers = [str(h or '').strip() for h in all_rows[0]]
            for row_vals in all_rows[1:]:
                if any(row_vals):
                    rows.append({headers[i]: (str(v or '').strip() if v is not None else '') for i, v in enumerate(row_vals) if i < len(headers)})

    created = 0
    updated = 0

    with db.transaction():
        c = db.get_company(company)
        if master_type == 'ledgers':
            for r in rows:
                name = r.get('Ledger Name') or r.get('name')
                if not name:
                    continue
                parent = r.get('Parent Group') or r.get('parent') or 'Sundry Debtors'
                record = {
                    'name': name,
                    'parent': parent,
                    'opening_balance': r.get('Opening Balance') or '0.00',
                    'drcr': r.get('Dr/Cr') or 'Dr',
                    'gstin': r.get('GSTIN') or '',
                    'state': r.get('State') or '',
                    'address': r.get('Address') or '',
                    'pincode': r.get('Pincode') or '',
                    'credit_period': r.get('Credit Days') or '',
                }
                is_new = db.save_master(company, 'ledger', record)
                if is_new: created += 1
                else: updated += 1
        elif master_type == 'items':
            for r in rows:
                name = r.get('Item Name') or r.get('name')
                if not name:
                    continue
                base_units = r.get('Base Unit') or r.get('base_units') or 'pcs'
                if base_units and base_units not in c.get('units', []):
                    db.save_master(company, 'unit', {'name': base_units})
                    c = db.get_company(company)
                record = {
                    'name': name,
                    'category': r.get('Stock Group') or 'Primary',
                    'base_units': base_units,
                    'alt_unit': r.get('Alt Unit') or '',
                    'conversion': r.get('Conversion Rate') or '',
                    'hsn': r.get('HSN Code') or '',
                    'gst_rate': r.get('GST Rate %') or '0',
                    'standard_purchase_rate': r.get('Standard Purchase Rate') or '',
                    'standard_selling_rate': r.get('Standard Selling Rate') or '',
                    'opening_qty': r.get('Opening Qty') or '0',
                    'opening_rate': r.get('Opening Rate') or '0',
                }
                is_new = db.save_master(company, 'stock_item', record)
                if is_new: created += 1
                else: updated += 1
        else:
            raise ValueError('Unsupported master type: ' + str(master_type))

    return {'status': 'success', 'created': created, 'updated': updated}
